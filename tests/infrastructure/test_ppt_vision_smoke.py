"""票 14 AC1：10 張 fixture 冒煙——真實 soffice 拆圖 → FakeVision 全鏈 → 注記 PDF。

真實視覺 API 屬 slow 冒煙（花錢＋上游依賴），以 FakeTranslator 取代；
本測試驗證「真 .pptx → soffice 渲染 → 每頁圖 → 組裝」管線（真實拆圖）。
LibreOffice 未裝時 skip（訊息含安裝指引）。
"""

from pathlib import Path

import pytest

from paper_kit.domain.vision_translation import VisionTranslation
from paper_kit.infrastructure.ppt_converter import LibreOfficeConverter, find_soffice
from paper_kit.infrastructure.ppt_vision_adapter import (
    PptVisionAdapter,
    PptVisionConfig,
)
from paper_kit.domain.translation_job import TranslationJob

pytestmark = pytest.mark.slow


@pytest.fixture(scope="module")
def soffice():
    soffice = find_soffice()
    if not soffice:
        pytest.skip("未偵測到 LibreOffice（winget install TheDocumentFoundation.LibreOffice）")
    return soffice


class FakeTranslator:
    def __init__(self, text: str):
        self._text = text

    def translate_image(self, image_path, target_lang="zh-TW"):
        return VisionTranslation(f"第 {image_path.name} 頁譯文：{self._text}", 10, 5)


def test_real_soffice_converts_ten_slide_deck(tmp_path, soffice):
    """10 張 fixture：python-pptx 造 10 頁 → soffice 渲染 → 10 張 PNG（依頁序）。"""
    from pptx import Presentation

    deck_path = tmp_path / "ten-slides.pptx"
    prs = Presentation()
    for i in range(10):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Slide {i + 1} Title"
        slide.placeholders[1].text = f"Content bullet {i + 1}: 論文翻譯測試"
    prs.save(str(deck_path))

    converter = LibreOfficeConverter(soffice=soffice)
    images = converter.convert_to_images(deck_path, tmp_path / "out")
    assert len(images) == 10
    assert [p.name for p in images] == [f"slide-{i:02d}.png" for i in range(1, 11)]
    assert all(p.stat().st_size > 0 for p in images)


def test_ppt_vision_pipeline_real_conversion_fake_vision(tmp_path, soffice):
    """全鏈：真拆圖（10 頁）＋假視覺 → 注記 PDF（10 頁、譯文可提取）。"""
    import pymupdf
    from pptx import Presentation

    deck_path = tmp_path / "deck.pptx"
    prs = Presentation()
    for i in range(10):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Title {i + 1}"
        slide.placeholders[1].text = f"Bullet {i + 1}"
    prs.save(str(deck_path))

    job = TranslationJob(job_id="smoke1", source_path=str(deck_path), target_lang="zh-TW")
    adapter = PptVisionAdapter(
        PptVisionConfig(api_key="FAKE-KEY"),
        converter=LibreOfficeConverter(soffice=soffice),
        translator=FakeTranslator("繁體中文譯文"),
    )
    result = adapter.translate(job)
    assert result.mono_path and Path(result.mono_path).exists()
    assert result.input_tokens == 100  # 10 頁 × 10
    doc = pymupdf.open(result.mono_path)
    assert doc.page_count == 10
    assert all(doc[i].get_images() for i in range(10))  # 每頁原圖（設計保留）
    assert any("繁體中文譯文" in doc[i].get_text() for i in range(10))
    doc.close()
